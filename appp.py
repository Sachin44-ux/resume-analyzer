import streamlit as st
import pickle
import docx
import PyPDF2
import re
import io
import os
import numpy as np
import pandas as pd
import matplotlib.pyplot as plt
import plotly.express as px
import plotly.graph_objects as go
import seaborn as sns
import smtplib
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText
from fpdf import FPDF

from typing import List, Tuple, Dict
from sklearn.metrics.pairwise import cosine_similarity
from sklearn.feature_extraction.text import ENGLISH_STOP_WORDS

# ==========================
# Load pre-trained artifacts
# ==========================
SVC_PATH   = r"C:\Users\ASUS\OneDrive\Desktop\Resume\clf.pkl"
TFIDF_PATH = r"C:\Users\ASUS\OneDrive\Desktop\Resume\tfidf.pkl"
LE_PATH    = r"C:\Users\ASUS\OneDrive\Desktop\Resume\encoder.pkl"

svc_model = pickle.load(open(SVC_PATH, "rb"))
tfidf     = pickle.load(open(TFIDF_PATH, "rb"))
le        = pickle.load(open(LE_PATH, "rb"))

# ==========================
# Utilities
# ==========================
PUNCT = re.escape("""!\"#$%&'()*+,-./:;<=>?@[\\]^_`{|}~""")
STOP  = set(ENGLISH_STOP_WORDS)

SKILL_TERMS = {"ui","ux","wireframe","prototype","figma","xd","sketch","photoshop","illustrator",
    "typography","layout","responsive","accessibility","wcag",
    "html","css","javascript","typescript","react","redux","angular","vue","node","express",
    "nextjs","vite","webpack","babel","tailwind","bootstrap","sass","less",
    "api","rest","graphql","jest","cypress","vitest","selenium",
    "seo","performance","lighthouse","pwa","webpack","git","github","gitlab",
    "mysql","postgres","mongodb","firebase","docker","kubernetes","ci","cd","aws","gcp","azure",
    "wordpress","shopify","webflow","drupal"
}

# --- Text Cleaning Functions ---
def clean_text_for_pdf(text: str) -> str:
    """Remove or replace characters that can't be encoded in latin-1"""
    if not text:
        return ""
    # Replace common problematic Unicode characters
    replacements = {
        '\u2014': '--',  # em dash
        '\u2013': '-',   # en dash
        '\u2018': "'",   # left single quote
        '\u2019': "'",   # right single quote
        '\u201c': '"',   # left double quote
        '\u201d': '"',   # right double quote
        '\u2022': '*',   # bullet point
        '\u2026': '...', # ellipsis
        '\u00a0': ' ',   # non-breaking space
        '\u200b': '',    # zero-width space
    }
    
    for unicode_char, replacement in replacements.items():
        text = text.replace(unicode_char, replacement)
    
    # Remove any other non-latin-1 characters
    try:
        text = text.encode('latin-1', 'replace').decode('latin-1')
    except (UnicodeEncodeError, UnicodeDecodeError):
        # If there are still issues, use a more aggressive approach
        text = text.encode('ascii', 'ignore').decode('ascii')
    
    return text

def clean_text(txt: str) -> str:
    if not txt:
        return ""
    txt = re.sub(r'http\S+\s', ' ', txt)
    txt = re.sub(r'\b(RT|cc)\b', ' ', txt)
    txt = re.sub(r'#\S+\s', ' ', txt)
    txt = re.sub(r'@\S+', ' ', txt)
    txt = re.sub(f'[{PUNCT}]', ' ', txt)
    txt = re.sub(r'[^\x00-\x7f]', ' ', txt)
    txt = re.sub(r'\s+', ' ', txt)
    return txt.strip().lower()

def pred_category(text: str) -> str:
    cleaned = clean_text(text)
    vec = tfidf.transform([cleaned]).toarray()
    pred = svc_model.predict(vec)
    return le.inverse_transform(pred)[0]

def vectorize_texts(texts: List[str]):
    cleaned = [clean_text(t) for t in texts]
    return tfidf.transform(cleaned)

def tokenize_words(txt: str) -> List[str]:
    txt = clean_text(txt)
    return [w for w in txt.split() if w not in STOP and len(w) > 2]

def jaccard(a: set, b: set) -> float:
    if not a and not b:
        return 0.0
    return len(a & b) / max(1, len(a | b))

def skills_overlap(jd_tokens: List[str], res_tokens: List[str]) -> float:
    a = set([t for t in jd_tokens if t in SKILL_TERMS])
    b = set([t for t in res_tokens if t in SKILL_TERMS])
    return jaccard(a, b)

def keyword_overlap(jd_tokens: List[str], res_tokens: List[str]) -> float:
    a = set(jd_tokens)
    b = set(res_tokens)
    return jaccard(a, b)

def chunk_text(text: str, chars: int = 1200) -> List[str]:
    text = text or ""
    text = text.strip()
    if len(text) <= chars:
        return [text]
    return [text[i:i+chars] for i in range(0, len(text), chars)]

def tfidf_cosine(a: str, b: str) -> float:
    A = vectorize_texts([a])
    B = vectorize_texts([b])
    return float(cosine_similarity(A, B).ravel()[0])

def chunk_max_cosine(jd: str, resume: str) -> float:
    jd_vec = vectorize_texts([jd])
    chunks = chunk_text(resume, 1200)
    if not chunks:
        return 0.0
    chunk_vecs = vectorize_texts(chunks)
    sims = cosine_similarity(jd_vec, chunk_vecs).ravel()
    return float(np.max(sims)) if len(sims) else 0.0

# --- File Extraction ---
def extract_text_from_pdf_bytes(b: bytes) -> str:
    try:
        reader = PyPDF2.PdfReader(io.BytesIO(b))
        return "\n".join([p.extract_text() or "" for p in reader.pages])
    except Exception:
        return ""

def extract_text_from_docx_bytes(b: bytes) -> str:
    try:
        document = docx.Document(io.BytesIO(b))
        return "\n".join(paragraph.text for paragraph in document.paragraphs)
    except Exception:
        return ""

def extract_text_from_txt_bytes(b: bytes) -> str:
    for enc in ("utf-8", "latin-1"):
        try:
            return b.decode(enc)
        except UnicodeDecodeError:
            continue
    return b.decode("utf-8", errors="ignore")

def extract_text_from_upload(uploaded_file) -> str:
    name = uploaded_file.name
    ext = os.path.splitext(name)[1].lower()
    data = uploaded_file.read()
    if ext == ".pdf":
        return extract_text_from_pdf_bytes(data)
    elif ext == ".docx":
        return extract_text_from_docx_bytes(data)
    elif ext == ".txt":
        return extract_text_from_txt_bytes(data)
    else:
        return ""

# ==========================
# Matching logic
# ==========================
def score_resume_vs_jd(jd_text: str, resume_text: str) -> Dict[str, float]:
    s_tfidf = tfidf_cosine(jd_text, resume_text)
    s_chunk = chunk_max_cosine(jd_text, resume_text)
    jd_tokens  = tokenize_words(jd_text)
    res_tokens = tokenize_words(resume_text)
    s_kw   = keyword_overlap(jd_tokens, res_tokens)
    s_sk   = skills_overlap(jd_tokens, res_tokens)
    composite = 0.6 * s_tfidf + 0.2 * s_chunk + 0.1 * s_kw + 0.1 * s_sk
    return {
        "cosine_tfidf": s_tfidf,
        "chunk_max": s_chunk,
        "keyword_overlap": s_kw,
        "skills_overlap": s_sk,
        "composite": composite
    }

def match_resumes_to_jd(resumes: List[Tuple[str, bytes, str]], jd_text: str, require_category_match: bool = False, min_similarity: float = 0.20, top_k: int = 10, ensure_topk_anyway: bool = True):
    jd_cat = pred_category(jd_text)
    rows = []
    for fname, fbytes, ftext in resumes:
        cat = pred_category(ftext) if ftext.strip() else "Unknown"
        scores = score_resume_vs_jd(jd_text, ftext)
        rows.append({"filename": fname, "bytes": fbytes, "text": ftext, "resume_category": cat, **scores})

    filtered = rows
    if require_category_match:
        _tmp = [r for r in rows if r["resume_category"] == jd_cat]
        filtered = _tmp if _tmp else rows

    _tmp = [r for r in filtered if r["composite"] >= min_similarity]
    if _tmp:
        filtered = _tmp
    else:
        _tmp2 = [r for r in filtered if r["composite"] >= max(0.10, min_similarity * 0.5)]
        filtered = _tmp2 if _tmp2 else filtered

    filtered.sort(key=lambda x: x["composite"], reverse=True)
    filtered = filtered[:top_k]

    if not filtered and ensure_topk_anyway:
        rows.sort(key=lambda x: x["composite"], reverse=True)
        filtered = rows[:min(top_k, len(rows))]

    return jd_cat, filtered

# ==========================
# Skill & Experience analysis
# ==========================
def analyze_resume_vs_jd(jd_tokens: List[str], res_tokens: List[str]):
    jd_skills = set([t for t in jd_tokens if t in SKILL_TERMS])
    res_skills = set([t for t in res_tokens if t in SKILL_TERMS])
    have = sorted(list(jd_skills & res_skills))
    missing = sorted(list(jd_skills - res_skills))
    return have, missing

def extract_experience_years(text: str) -> int:
    match = re.search(r'(\d{1,2})\+?\s*(?:years|yrs)', text.lower())
    if match:
        return int(match.group(1))
    return 0

# ==========================
# Resume Improvement (NLP Suggestions)
# ==========================
ACTION_VERBS = {
    "optimized","developed","designed","created","led","managed",
    "implemented","built","analyzed","trained","deployed",
    "improved","achieved","automated","engineered","delivered"
}

def extract_sections(text: str) -> Dict[str, str]:
    sections = {
        "summary": "",
        "experience": "",
        "projects": "",
        "skills": "",
        "education": "",
    }

    text_low = text.lower()

    patterns = {
        "summary": r"(summary|profile|objective)[:\n](.*?)(?=\n[a-z ]+[:\n]|$)",
        "experience": r"(experience|work experience|employment)[:\n](.*?)(?=\n[a-z ]+[:\n]|$)",
        "projects": r"(projects|project experience)[:\n](.*?)(?=\n[a-z ]+[:\n]|$)",
        "skills": r"(skills|technical skills)[:\n](.*?)(?=\n[a-z ]+[:\n]|$)",
        "education": r"(education|academic background)[:\n](.*?)(?=\n[a-z ]+[:\n]|$)",
    }

    for key, pattern in patterns.items():
        m = re.search(pattern, text_low, re.DOTALL)
        if m:
            sections[key] = m.group(2).strip()

    return sections


def generate_resume_improvements(resume_text: str, jd_text: str) -> Dict[str, List[str]]:
    suggestions = {
        "missing_keywords": [],
        "missing_skills": [],
        "weak_summary": [],
        "weak_experience": [],
        "bullet_issues": [],
        "general_advice": []
    }

    jd_tokens = tokenize_words(jd_text)
    res_tokens = tokenize_words(resume_text)

    jd_top_words = [w for w in jd_tokens if len(w) > 3]
    if len(jd_top_words) > 0:
        jd_freq = pd.Series(jd_top_words).value_counts().head(20).index.tolist()
    else:
        jd_freq = []

    missing_kw = [w for w in jd_freq if w not in res_tokens]

    if missing_kw:
        suggestions["missing_keywords"] = missing_kw

    jd_skills = {t for t in jd_tokens if t in SKILL_TERMS}
    res_skills = {t for t in res_tokens if t in SKILL_TERMS}

    missing_sk = sorted(list(jd_skills - res_skills))

    if missing_sk:
        suggestions["missing_skills"] = missing_sk

    sections = extract_sections(resume_text)

    summary = sections["summary"]
    if not summary or len(summary.split()) < 25:
        suggestions["weak_summary"].append(
            "Summary is too short — add 2–3 lines highlighting your strongest achievements, tech stack, and domain expertise."
        )

    experience = sections["experience"]
    if experience and len(experience.split()) < 40:
        suggestions["weak_experience"].append(
            "Experience section seems short — add detailed bullet points describing impact, metrics, and technologies used."
        )

    bullets = re.findall(r"(?:^|\n)\s*(?:•|-|\*|\u2022)\s+", resume_text)
    if len(bullets) < 4:
        suggestions["bullet_issues"].append(
            "Add more bullet points — each position should have 3–5 achievement-based bullets."
        )

    if not any(v in resume_text.lower() for v in ACTION_VERBS):
        suggestions["bullet_issues"].append(
            "Use strong action verbs such as 'developed', 'optimized', 'implemented' to describe achievements."
        )

    suggestions["general_advice"] = [
        "Add quantifiable metrics: e.g., 'Improved efficiency by 30%'",
        "Ensure formatting is consistent (fonts, headings, bullet alignment)",
        "Avoid long paragraphs — use short bullet points",
        "Keep resume to 1 page if experience < 4 years"
    ]

    return suggestions

# ==========================
# ATS Score Calculation (0–100)
# ==========================
def calculate_ats_score(resume_text: str, jd_text: str) -> int:
    res_tokens = tokenize_words(resume_text)
    jd_tokens = tokenize_words(jd_text)

    jd_keywords = set([w for w in jd_tokens if len(w) > 3])
    matched_keywords = len([w for w in res_tokens if w in jd_keywords])
    keyword_score = min(1, matched_keywords / (len(jd_keywords) + 1))

    jd_skills = {w for w in jd_tokens if w in SKILL_TERMS}
    res_skills = {w for w in res_tokens if w in SKILL_TERMS}
    matched_skills = len(jd_skills & res_skills)
    skill_score = min(1, matched_skills / (len(jd_skills) + 1))

    ACTIONS = {
        "developed", "designed", "implemented", "created", "built", "led",
        "optimized", "automated", "engineered", "managed", "improved"
    }
    action_score = 1 if any(a in resume_text.lower() for a in ACTIONS) else 0

    formatting_score = 1 if ("•" in resume_text or "-" in resume_text or "*" in resume_text) else 0

    final_score = (
        keyword_score * 40 +
        skill_score * 40 +
        action_score * 10 +
        formatting_score * 10
    )

    return int(final_score)

# ==========================
# PDF Report Generation (FIXED)
# ==========================
def generate_report_pdf(filename: str, ats_score: int, suggestions: Dict[str, List[str]], have: List[str], missing: List[str], jd_text: str, resume_text: str) -> bytes:
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    pdf.set_font("Arial", 'B', 16)
    pdf.cell(0, 10, "Resume Analysis Report", ln=True, align='C')
    pdf.ln(4)

    pdf.set_font("Arial", size=12)
    
    # Clean all text before adding to PDF
    clean_filename = clean_text_for_pdf(filename)
    clean_have = clean_text_for_pdf(', '.join(have) if have else 'None')
    clean_missing = clean_text_for_pdf(', '.join(missing) if missing else 'None')
    clean_jd_text = clean_text_for_pdf(jd_text[:1000])
    clean_resume_text = clean_text_for_pdf(resume_text[:1200])
    
    pdf.cell(40, 8, f"Filename: {clean_filename}", ln=True)
    pdf.cell(40, 8, f"ATS Score: {ats_score}/100", ln=True)
    pdf.cell(40, 8, f"Matched Skills: {clean_have}", ln=True)
    pdf.cell(40, 8, f"Missing Skills: {clean_missing}", ln=True)
    pdf.ln(6)

    pdf.set_font("Arial", 'B', 12)
    pdf.cell(0, 8, "Key Suggestions:", ln=True)
    pdf.set_font("Arial", size=11)

    # Clean and add suggestions
    if suggestions.get("missing_keywords"):
        clean_missing_kw = clean_text_for_pdf(", ".join(suggestions["missing_keywords"]))
        pdf.multi_cell(0, 6, "Missing Keywords: " + clean_missing_kw)
    if suggestions.get("missing_skills"):
        clean_missing_skills = clean_text_for_pdf(", ".join(suggestions["missing_skills"]))
        pdf.multi_cell(0, 6, "Missing Skills: " + clean_missing_skills)

    if suggestions.get("weak_summary"):
        for l in suggestions.get("weak_summary"):
            clean_l = clean_text_for_pdf(l)
            pdf.multi_cell(0, 6, "- " + clean_l)
    if suggestions.get("weak_experience"):
        for l in suggestions.get("weak_experience"):
            clean_l = clean_text_for_pdf(l)
            pdf.multi_cell(0, 6, "- " + clean_l)
    if suggestions.get("bullet_issues"):
        for l in suggestions.get("bullet_issues"):
            clean_l = clean_text_for_pdf(l)
            pdf.multi_cell(0, 6, "- " + clean_l)

    pdf.ln(6)
    pdf.set_font("Arial", 'B', 12)
    pdf.cell(0, 8, "General Advice:", ln=True)
    pdf.set_font("Arial", size=11)
    for tip in suggestions.get("general_advice", []):
        clean_tip = clean_text_for_pdf(tip)
        pdf.multi_cell(0, 6, "- " + clean_tip)

    pdf.ln(6)
    pdf.set_font("Arial", 'B', 12)
    pdf.cell(0, 8, "Job Description (snippet):", ln=True)
    pdf.set_font("Arial", size=10)
    pdf.multi_cell(0, 6, clean_jd_text)

    pdf.ln(6)
    pdf.set_font("Arial", 'B', 12)
    pdf.cell(0, 8, "Resume (snippet):", ln=True)
    pdf.set_font("Arial", size=10)
    pdf.multi_cell(0, 6, clean_resume_text)

    # FIX: Use the correct method to output to bytes
    return pdf.output(dest='S').encode('latin-1')

# ==========================
# Streamlit UI
# ==========================

def main():
    st.set_page_config(
        page_title="AI Resume Ranker | Multi-role",
        page_icon="🧠",
        layout="wide",
    )

    st.markdown("""
    <style>
        .big-title { font-size: 32px !important; font-weight: 700 !important; color: #1F4172 !important; }
        .sub-heading { font-size: 20px !important; font-weight: 600 !important; margin-top: 6px; color: #26355D !important; }
        .section-box { padding: 14px; border-radius: 10px; background: #f8f9fa; border: 1px solid #E0E0E0; margin-bottom: 14px; }
    </style>
    """, unsafe_allow_html=True)

    st.markdown("<div class='big-title'>🧠 AI Resume Ranker — Students & Recruiters</div>", unsafe_allow_html=True)
    st.write("Choose your role: analyze a single resume (Students/Working Professional) or run a recruiter workflow for multiple candidates.")
    st.write("---")

    role_tabs = st.tabs(["Students / Working Professional", "Recruiter"])

    # --------------------------
    # Students / Working Professional
    # --------------------------
    with role_tabs[0]:
        st.markdown("<div class='sub-heading'>👩‍🎓 Resume Analyzer (Individual)</div>", unsafe_allow_html=True)
        st.write("Upload your CV (PDF/DOCX/TXT). Optionally paste a JD to evaluate your resume against a role. If you don't provide a JD, a generic role profile will be used.")

        uploaded = st.file_uploader("Upload your Resume", type=["pdf","docx","txt"], key="student_upload")
        jd_text_student = st.text_area("Paste Job Description (optional)", height=180, key="student_jd")

        analyze_btn = st.button("Analyze Resume — Student", key="student_analyze")

        if analyze_btn:
            if not uploaded:
                st.error("Please upload your resume file.")
            else:
                with st.spinner("Analyzing your resume..."):
                    resume_text = extract_text_from_upload(uploaded)
                    if not jd_text_student.strip():
                        # build a simple generic JD from SKILL_TERMS
                        jd_text_student = "JOB: Generic Software Role\nRequired: " + ", ".join(list(SKILL_TERMS)[:40])

                    ats = calculate_ats_score(resume_text, jd_text_student)
                    st.metric("ATS Score", f"{ats}/100")

                    suggestions = generate_resume_improvements(resume_text, jd_text_student)
                    jd_tokens = tokenize_words(jd_text_student)
                    res_tokens = tokenize_words(resume_text)
                    have, missing = analyze_resume_vs_jd(jd_tokens, res_tokens)

                    if suggestions.get("missing_keywords"):
                        st.warning("Missing Keywords: " + ", ".join(suggestions["missing_keywords"]))
                    if suggestions.get("missing_skills"):
                        st.error("Missing Required Skills: " + ", ".join(suggestions["missing_skills"]))

                    st.subheader("Improvement Suggestions")
                    for section_title, items in suggestions.items():
                        if items and section_title not in ["missing_keywords", "missing_skills"]:
                            st.write("**" + section_title.replace('_',' ').title() + ":**")
                            for it in items:
                                st.write("•", it)

                    st.subheader("Matched / Missing Skills")
                    st.write("Matched Skills:", ", ".join(have) if have else "None")
                    st.write("Missing Skills:", ", ".join(missing) if missing else "None")

                    # Download analysis as PDF
                    try:
                        pdf_bytes = generate_report_pdf(uploaded.name, ats, suggestions, have, missing, jd_text_student, resume_text)
                        st.download_button(
                            "Download Analysis PDF", 
                            data=pdf_bytes, 
                            file_name=f"analysis_{os.path.splitext(uploaded.name)[0]}.pdf", 
                            mime="application/pdf",
                            key="student_pdf"
                        )
                    except Exception as e:
                        st.error(f"Error generating PDF: {str(e)}")

    # --------------------------
    # Recruiter (original deeper flow)
    # --------------------------
    with role_tabs[1]:
        st.markdown("<div class='sub-heading'>💼 Recruiter Dashboard (Batch)</div>", unsafe_allow_html=True)
        st.write("Upload multiple resumes and a Job Description. This tab preserves the original recruiter features: ranking, visualization, and CSV export.")

        # Sidebar-like controls placed inline for convenience
        require_match = st.checkbox("Require category match", value=False, key="req_match")
        min_sim = st.slider("Min composite score", 0.0, 1.0, 0.20, 0.01, key="min_sim")
        top_k = st.number_input("Top K Results", 1, 50, 10, key="top_k")
        ensure_any = st.checkbox("Always return Top-K", value=True, key="ensure_any")

        resumes = st.file_uploader(
            "Upload Resumes (PDF/DOCX/TXT)",
            type=["pdf", "docx", "txt"],
            accept_multiple_files=True,
            key="recruiter_uploads"
        )

        jd_tab, file_tab = st.tabs(["✍️ Paste JD", "📄 Upload JD"])
        jd_text_manual = ""
        jd_file = None
        with jd_tab:
            jd_text_manual = st.text_area("Paste Job Description", height=180, placeholder="Enter JD here...", key="jd_paste")
        with file_tab:
            jd_file = st.file_uploader("Upload JD", type=["pdf", "docx", "txt"], key="jd_upload")

        process = st.button("🚀 Run Resume Ranking — Recruiter", key="recruiter_process")

        if process:
            if not resumes:
                st.error("Please upload at least one resume.")
            else:
                jd_text = jd_text_manual.strip()
                if not jd_text and jd_file:
                    jd_text = extract_text_from_upload(jd_file)

                if not jd_text:
                    st.error("Please provide a job description.")
                else:
                    resume_triplets = []
                    with st.spinner("Extracting text from resumes..."):
                        for f in resumes:
                            data = f.read()
                            ext = os.path.splitext(f.name)[1].lower()

                            if ext == ".pdf":
                                text = extract_text_from_pdf_bytes(data)
                            elif ext == ".docx":
                                text = extract_text_from_docx_bytes(data)
                            else:
                                text = extract_text_from_txt_bytes(data)

                            resume_triplets.append((f.name, data, text))

                    with st.spinner("Analyzing and ranking resumes..."):
                        jd_category, matches = match_resumes_to_jd(
                            resumes=resume_triplets,
                            jd_text=jd_text,
                            require_category_match=require_match,
                            min_similarity=min_sim,
                            top_k=top_k,
                            ensure_topk_anyway=ensure_any
                        )

                    st.success(f"JD Category Predicted: **{jd_category}**")

                    if not matches:
                        st.warning("No matching resumes found.")
                    else:
                        jd_tokens = tokenize_words(jd_text)

                        table_data = []
                        for row in matches:
                            res_tokens = tokenize_words(row["text"])
                            have, missing = analyze_resume_vs_jd(jd_tokens, res_tokens)
                            exp_years = extract_experience_years(row["text"])
                            table_data.append({
                                "Filename": row["filename"],
                                "Category": row["resume_category"],
                                "Composite Score": round(row["composite"], 3),
                                "TF-IDF": round(row["cosine_tfidf"], 3),
                                "Skill Match": round(row["skills_overlap"], 3),
                                "Experience (yrs)": exp_years,
                                "Matched Skills": ", ".join(have),
                                "Missing Skills": ", ".join(missing)
                            })

                        df = pd.DataFrame(table_data)
                        st.dataframe(df, use_container_width=True)

                        st.download_button("⬇️ Download Results CSV", df.to_csv(index=False), "results.csv", key="csv_download")

                        vis1, vis2, vis3 = st.tabs(["🏆 Ranking", "📌 Skill Radar", "📉 Experience"])

                        with vis1:
                            fig = px.bar(df, y="Filename", x="Composite Score", orientation="h", color="Composite Score")
                            st.plotly_chart(fig, use_container_width=True)

                        with vis2:
                            if len(df) > 0:
                                fig = go.Figure()
                                for _, r in df.iterrows():
                                    fig.add_trace(go.Scatterpolar(
                                        r=[r["TF-IDF"], r["Skill Match"], min(r["Experience (yrs)"], 10)],  # Cap experience for better visualization
                                        theta=["TF-IDF", "Skill Match", "Experience"],
                                        fill="toself",
                                        name=r["Filename"][:30]  # Truncate long filenames
                                    ))
                                st.plotly_chart(fig, use_container_width=True)
                            else:
                                st.info("No data available for radar chart")

                        with vis3:
                            if len(df) > 0:
                                fig, ax = plt.subplots()
                                ax.hist(df["Experience (yrs)"], bins=range(0, max(df["Experience (yrs)"].max() + 2, 2)))
                                ax.set_xlabel("Experience (years)")
                                ax.set_ylabel("Number of Candidates")
                                st.pyplot(fig)
                            else:
                                st.info("No data available for experience histogram")

                        st.markdown("<div class='sub-heading'>🔍 Candidate Breakdown</div>", unsafe_allow_html=True)

                        for i, row in enumerate(matches, 1):
                            with st.expander(f"{i}. {row['filename']} — Score: {row['composite']:.3f}"):
                                st.text_area(f"Resume Preview {i}", row["text"][:1200], height=190, key=f"preview_{i}")

                                ats_score = calculate_ats_score(row["text"], jd_text)
                                st.metric("ATS Score", f"{ats_score}/100")

                                suggestions = generate_resume_improvements(row["text"], jd_text)

                                if suggestions["missing_keywords"]:
                                    st.warning("Missing Keywords: " + ", ".join(suggestions["missing_keywords"]))

                                if suggestions["missing_skills"]:
                                    st.error("Missing Required Skills: " + ", ".join(suggestions["missing_skills"]))

                                for section_title, items in suggestions.items():
                                    if items and section_title not in ["missing_keywords", "missing_skills", "general_advice"]:
                                        st.info(f"**{section_title.replace('_',' ').title()}:**")
                                        for it in items:
                                            st.write("•", it)

                                st.success("General Advice")
                                for tip in suggestions["general_advice"]:
                                    st.write("•", tip)

                                # per-candidate PDF download
                                jd_tokens_local = tokenize_words(jd_text)
                                res_tokens_local = tokenize_words(row["text"])
                                have_local, missing_local = analyze_resume_vs_jd(jd_tokens_local, res_tokens_local)
                                try:
                                    pdf_bytes = generate_report_pdf(row["filename"], ats_score, suggestions, have_local, missing_local, jd_text, row["text"])
                                    st.download_button(
                                        f"Download {row['filename']} Report", 
                                        data=pdf_bytes, 
                                        file_name=f"analysis_{os.path.splitext(row['filename'])[0]}.pdf", 
                                        mime="application/pdf",
                                        key=f"pdf_{i}"
                                    )
                                except Exception as e:
                                    st.error(f"Error generating PDF for {row['filename']}: {str(e)}")


if __name__ == "__main__":
    main()